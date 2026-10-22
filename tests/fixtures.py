"""Synthetic test documents, generated on the fly so no real document ever enters the repo."""
from __future__ import annotations

import io
from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def policy_pdf(path: Path, token: str = "") -> Path:
    """Two pages: a product summary with a ruled premium table, then a page about exclusions."""
    styles = getSampleStyleSheet()
    table = Table([
        ["Age band", "Annual premium (S$)", "Sum assured (S$)"],
        ["18 - 30", "320", "100,000"],
        ["31 - 45", "540", "100,000"],
        ["46 - 60", "1,180", "100,000"],
    ])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    story = [
        Paragraph("Singlife Shield Saver Product Summary", styles["Title"]),
        Paragraph("Version 3. Effective date: 1 April 2026.", styles["Normal"]),
        Paragraph(f"This plan covers hospitalisation and surgical expenses. {token}", styles["Normal"]),
        Spacer(1, 12),
        Paragraph("Figure 1: Premium by age band", styles["Normal"]),
        table,
        PageBreak(),
        Paragraph("Exclusions", styles["Heading2"]),
        Paragraph("Pre-existing conditions are excluded during the first 12 months. "
                  "Cosmetic surgery is not covered.", styles["Normal"]),
    ]
    SimpleDocTemplate(str(path), pagesize=A4).build(story)
    return path


def scanned_pdf(path: Path, text: str) -> Path:
    """A PDF whose only page is an image of text: no text layer, so it needs OCR."""
    from PIL import Image as PILImage, ImageDraw, ImageFont

    img = PILImage.new("RGB", (1700, 400), "white")
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 48)
    except OSError:
        font = ImageFont.load_default(size=48)
    ImageDraw.Draw(img).text((60, 150), text, fill="black", font=font)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    SimpleDocTemplate(str(path), pagesize=A4).build([Image(buf, width=500, height=118)])
    return path


def policy_docx(path: Path, token: str = "") -> Path:
    import docx

    d = docx.Document()
    d.add_heading("Manulife Brochure", 1)
    d.add_paragraph(f"Critical illness cover for early and major stages. {token}")
    t = d.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("Benefit", "Payout"), ("Early stage", "25%"), ("Major stage", "100%")]):
        t.cell(r, 0).text, t.cell(r, 1).text = a, b
    d.save(path)
    return path


def rates_xlsx(path: Path, token: str = "") -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rates"
    ws.append(["Plan", "Term", "Rate"])
    ws.append(["Endowment 10", "10 years", f"2.08% {token}"])
    ws.append(["Endowment 20", "20 years", "2.50%"])
    wb.save(path)
    return path
