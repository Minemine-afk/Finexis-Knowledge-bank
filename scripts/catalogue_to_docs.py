"""Turn the static Policy Desk catalogue (index.html) into one Word document per insurer, so it can
be uploaded to the knowledge bank like any other document.

    python -m scripts.catalogue_to_docs samples/catalogue

Estimated fields (the record's `est` list) are labelled as estimates in the text so a search hit
can never pass them off as the insurer's own terms. Output goes under samples/, which git ignores.
"""
from __future__ import annotations

import json
import re
import subprocess
import sys
from collections import defaultdict
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
FIELDS = [("category", "Category"), ("type", "Product type"), ("status", "Distribution status"),
          ("premiumBasis", "Premium basis"), ("coverageTerm", "Coverage term"), ("entryAge", "Entry age"),
          ("sumAssured", "Sum assured / benefit"), ("uw", "Underwriting route"),
          ("funding", "Funding"), ("currency", "Currency"), ("needs", "Client needs"), ("segments", "Segments")]


def load_products() -> list[dict]:
    html = (ROOT / "index.html").read_text(encoding="utf-8")
    m = re.search(r"const PRODUCTS = (\[.*?\n\]);", html, re.S)
    if not m:
        raise SystemExit("PRODUCTS not found in index.html")
    js = f"console.log(JSON.stringify({m.group(1)}))"
    return json.loads(subprocess.run(["node", "-e", js], check=True, capture_output=True, text=True).stdout)


def write_insurer(insurer: str, products: list[dict], out: Path) -> Path:
    doc = Document()
    doc.add_heading(f"{insurer} product summary", 0)
    sourced = max((p.get("sourced") or "" for p in products), default="")
    doc.add_paragraph(f"Issuer: {insurer}. Document type: product summary (demo catalogue). "
                      f"Information taken from the insurer's public pages on {sourced}. "
                      "Fields marked (estimate) are NOT from the insurer and must be checked against the "
                      "official product summary before use.")
    for p in products:
        doc.add_heading(p["name"], 1)
        doc.add_paragraph(p.get("summary", ""))
        est = set(p.get("est", []))
        for key, label in FIELDS:
            value = p.get(key)
            if value:
                text = ", ".join(value) if isinstance(value, list) else str(value)
                doc.add_paragraph(f"{label}: {text}" + (" (estimate)" if key in est else ""))
        for heading, key in (("Highlights", "features"), ("Adviser notes", "notes")):
            if p.get(key):
                doc.add_paragraph(f"{heading}:")
                for line in p[key]:
                    doc.add_paragraph(line, style="List Bullet")
        if p.get("source"):
            doc.add_paragraph(f"Source: {p['source']}")
    path = out / f"{re.sub(r'[^A-Za-z0-9]+', '_', insurer).strip('_')}_product_summary.docx"
    doc.save(path)
    return path


def main() -> None:
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "samples/catalogue")
    out.mkdir(parents=True, exist_ok=True)
    by_insurer = defaultdict(list)
    for p in load_products():
        by_insurer[p["insurer"]].append(p)
    for insurer, products in sorted(by_insurer.items()):
        print(write_insurer(insurer, products, out), f"({len(products)} products)")


if __name__ == "__main__":
    main()
